"""Kiểm thử: inputs."""

import tempfile
from unittest.mock import MagicMock, patch

from django.contrib.auth.models import User
from django.core.cache import cache
from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import TestCase, override_settings
from django.urls import reverse

from .forms import SettingsForm, SummaryRequestForm
from .readers import _is_private_ip, _resolve_and_validate, extract_text


class FormValidationTests(TestCase):
    def test_settings_form_valid(self):
        form = SettingsForm(data={"default_summary_ratio": 0.3, "gemini_api_key": ""})
        self.assertTrue(form.is_valid())

    def test_settings_form_negative_ratio(self):
        form = SettingsForm(data={"default_summary_ratio": -1, "gemini_api_key": ""})
        self.assertFalse(form.is_valid())

    def test_settings_form_ratio_too_high(self):
        form = SettingsForm(data={"default_summary_ratio": 2.0, "gemini_api_key": ""})
        self.assertFalse(form.is_valid())

    def test_settings_form_long_api_key(self):
        form = SettingsForm(data={"default_summary_ratio": 0.2, "gemini_api_key": "k" * 500})
        self.assertFalse(form.is_valid())

    def test_summary_form_missing_source_type(self):
        form = SummaryRequestForm(data={})
        self.assertFalse(form.is_valid())


class UrlExtractionTests(TestCase):
    def _mock_session_get(self, mock_get):
        mock_get.return_value.status_code = 200
        mock_get.return_value.headers = {"Content-Type": "text/html; charset=utf-8"}

    @patch("summaries.readers._get_http_session")
    def test_extract_text_from_url(self, mock_session):
        sess = mock_session.return_value
        sess.get.return_value.status_code = 200
        sess.get.return_value.headers = {"Content-Type": "text/html; charset=utf-8"}
        sess.get.return_value.text = (
            "<html><body><p>Hello world. This is a test page.</p></body></html>"
        )
        result = extract_text("https://example.com")
        self.assertIn("Hello world", result)

    @patch("summaries.readers._get_http_session")
    def test_extract_text_from_url_removes_scripts(self, mock_session):
        sess = mock_session.return_value
        sess.get.return_value.status_code = 200
        sess.get.return_value.headers = {"Content-Type": "text/html; charset=utf-8"}
        sess.get.return_value.text = (
            "<html><head><script>alert('xss')</script></head>"
            "<body><p>Main content here.</p></body></html>"
        )
        result = extract_text("https://example.com")
        self.assertIn("Main content", result)
        self.assertNotIn("alert", result)

    @patch("summaries.readers._get_http_session")
    def test_extract_text_from_url_invalid_content_type(self, mock_session):
        sess = mock_session.return_value
        sess.get.return_value.status_code = 200
        sess.get.return_value.headers = {"Content-Type": "application/pdf"}
        sess.get.return_value.text = "not html"
        with self.assertRaises(ValueError):
            extract_text("https://example.com/file.pdf")

    @patch("summaries.readers._get_http_session")
    def test_extract_text_from_url_raises_on_timeout(self, mock_session):
        from requests.exceptions import Timeout as RequestsTimeout

        sess = mock_session.return_value
        sess.get.side_effect = RequestsTimeout("Timeout")
        with self.assertRaises(ValueError):
            extract_text("https://example.com")

    def test_extract_text_from_url_invalid_scheme(self):
        with self.assertRaises((ValueError, FileNotFoundError)):
            extract_text("ftp://example.com")

    def test_extract_text_from_url_no_netloc(self):
        with self.assertRaises(ValueError):
            extract_text("http://")


class SsrfProtectionTests(TestCase):
    def test_is_private_ip_blocks_loopback(self):
        self.assertTrue(_is_private_ip("127.0.0.1"))
        self.assertTrue(_is_private_ip("::1"))

    def test_is_private_ip_blocks_private_ranges(self):
        self.assertTrue(_is_private_ip("10.0.0.5"))
        self.assertTrue(_is_private_ip("192.168.1.1"))
        self.assertTrue(_is_private_ip("172.16.0.1"))
        self.assertTrue(_is_private_ip("169.254.169.254"))
        self.assertTrue(_is_private_ip("fd00::1"))

    def test_is_private_ip_allows_public(self):
        self.assertFalse(_is_private_ip("8.8.8.8"))
        self.assertFalse(_is_private_ip("93.184.216.34"))

    @patch("summaries.readers.socket.getaddrinfo")
    def test_resolve_and_validate_blocks_private(self, mock_getaddrinfo):
        mock_getaddrinfo.return_value = [(None, None, None, None, ("127.0.0.1", 80))]
        with self.assertRaises(ValueError):
            _resolve_and_validate("localhost")

    @patch("summaries.readers.socket.getaddrinfo")
    def test_resolve_and_validate_allows_public(self, mock_getaddrinfo):
        mock_getaddrinfo.return_value = [(None, None, None, None, ("8.8.8.8", 80))]
        _resolve_and_validate("example.com")

    @patch("summaries.readers.socket.getaddrinfo")
    def test_extract_url_blocks_private_hostname(self, mock_getaddrinfo):
        mock_getaddrinfo.return_value = [(None, None, None, None, ("127.0.0.1", 80))]
        with self.assertRaises(ValueError):
            extract_text("http://localhost/secret")


class UrlSourceFlowTests(TestCase):
    def setUp(self):
        self.user = User.objects.create_user(username="url-test", password="secret123")

        self.client.login(username="url-test", password="secret123")

    def tearDown(self):
        cache.clear()

    @patch("summaries.readers._get_http_session")
    def test_create_summary_from_url_view(self, mock_session):
        sess = mock_session.return_value
        sess.get.return_value.status_code = 200
        sess.get.return_value.headers = {"Content-Type": "text/html; charset=utf-8"}
        sess.get.return_value.text = (
            "<html><body><p>First useful sentence. "
            "Second useful sentence. Third sentence.</p></body></html>"
        )
        response = self.client.post(
            reverse("create_summary"),
            {
                "source_type": "url",
                "source_url": "https://example.com/article",
                "method": "textrank",
                "ratio": 0.3,
            },
        )
        self.assertEqual(response.status_code, 200)
        data = response.json()
        self.assertTrue(data["ok"])
        self.assertEqual(data["data"]["source_type"], "url")

    def test_create_summary_url_requires_url(self):
        response = self.client.post(
            reverse("create_summary"),
            {
                "source_type": "url",
                "method": "textrank",
                "ratio": 0.3,
            },
        )
        self.assertEqual(response.status_code, 400)


@override_settings(MEDIA_ROOT=tempfile.mkdtemp(), RATE_LIMIT_SECONDS=0)
class FileUploadTests(TestCase):
    def setUp(self):
        from django.core.cache import cache

        cache.clear()
        self.user = User.objects.create_user(username="uploader", password="secret123")

        self.client.login(username="uploader", password="secret123")

    def test_upload_txt_file(self):
        text_content = b"This is a test document. It has multiple sentences. We need enough text."
        uploaded = SimpleUploadedFile("test.txt", text_content, content_type="text/plain")
        response = self.client.post(
            reverse("create_summary"),
            {
                "source_type": "file",
                "method": "textrank",
                "upload": uploaded,
                "ratio": 0.3,
            },
        )
        self.assertEqual(response.status_code, 200)
        data = response.json()
        self.assertTrue(data["ok"])

    def test_upload_file_too_large(self):
        big_content = b"x" * (10 * 1024 * 1024 + 1)
        uploaded = SimpleUploadedFile("big.txt", big_content, content_type="text/plain")
        response = self.client.post(
            reverse("create_summary"),
            {
                "source_type": "file",
                "method": "textrank",
                "upload": uploaded,
                "ratio": 0.3,
            },
        )
        self.assertEqual(response.status_code, 400)

    def test_upload_no_file_sent(self):
        response = self.client.post(
            reverse("create_summary"),
            {
                "source_type": "file",
                "method": "textrank",
                "ratio": 0.3,
            },
        )
        self.assertEqual(response.status_code, 400)

    def test_upload_unsupported_format(self):
        uploaded = SimpleUploadedFile(
            "test.exe", b"fake content", content_type="application/octet-stream"
        )
        response = self.client.post(
            reverse("create_summary"),
            {
                "source_type": "file",
                "method": "textrank",
                "upload": uploaded,
                "ratio": 0.3,
            },
        )
        self.assertEqual(response.status_code, 400)

    def test_upload_empty_content_extraction(self):
        uploaded = SimpleUploadedFile("empty.txt", b"   ", content_type="text/plain")
        response = self.client.post(
            reverse("create_summary"),
            {
                "source_type": "file",
                "method": "textrank",
                "upload": uploaded,
                "ratio": 0.3,
            },
        )
        self.assertEqual(response.status_code, 400)
        self.assertIn("Không thể trích xuất", response.json()["message"])


@override_settings(RATE_LIMIT_SECONDS=0)
class FileExtractionTests(TestCase):
    def setUp(self):
        self.user = User.objects.create_user(username="extract-test", password="secret123")

        self.client.login(username="extract-test", password="secret123")

    def tearDown(self):
        cache.clear()

    @patch("summaries.readers._extract_text_from_txt")
    def test_txt_file_extraction(self, mock_extract):
        mock_extract.return_value = (
            "This is extracted text from a txt file. It has multiple sentences."
        )
        uploaded = SimpleUploadedFile("test.txt", b"ignored", content_type="text/plain")
        response = self.client.post(
            reverse("create_summary"),
            {
                "source_type": "file",
                "method": "textrank",
                "upload": uploaded,
                "ratio": 0.3,
            },
        )
        self.assertEqual(response.status_code, 200)
        self.assertTrue(response.json()["ok"])
        mock_extract.assert_called_once()

    @patch("summaries.readers._extract_text_from_docx")
    def test_docx_file_extraction(self, mock_extract):
        mock_extract.return_value = "Extracted content from a DOCX file."
        uploaded = SimpleUploadedFile(
            "test.docx",
            b"ignored",
            content_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        )
        response = self.client.post(
            reverse("create_summary"),
            {
                "source_type": "file",
                "method": "textrank",
                "upload": uploaded,
                "ratio": 0.3,
            },
        )
        self.assertEqual(response.status_code, 200)
        mock_extract.assert_called_once()

    @patch("summaries.readers._extract_text_from_pdf")
    def test_pdf_file_extraction(self, mock_extract):
        mock_extract.return_value = "Extracted content from a PDF file."
        uploaded = SimpleUploadedFile("test.pdf", b"ignored", content_type="application/pdf")
        response = self.client.post(
            reverse("create_summary"),
            {
                "source_type": "file",
                "method": "textrank",
                "upload": uploaded,
                "ratio": 0.3,
            },
        )
        self.assertEqual(response.status_code, 200)
        mock_extract.assert_called_once()

    @patch("summaries.readers._extract_text_from_epub")
    def test_epub_file_extraction(self, mock_extract):
        mock_extract.return_value = "Extracted content from an EPUB file."
        uploaded = SimpleUploadedFile("test.epub", b"ignored", content_type="application/epub+zip")
        response = self.client.post(
            reverse("create_summary"),
            {
                "source_type": "file",
                "method": "textrank",
                "upload": uploaded,
                "ratio": 0.3,
            },
        )
        self.assertEqual(response.status_code, 200)
        mock_extract.assert_called_once()


class ReaderUnitTests(TestCase):
    def test_extract_text_from_txt_utf8(self):
        import pathlib
        import tempfile as _tf

        from .readers import _extract_text_from_txt

        with _tf.NamedTemporaryFile("w", suffix=".txt", delete=False, encoding="utf-8") as f:
            f.write("Xin chào thế giới. Đây là nội dung UTF-8.")
            p = pathlib.Path(f.name)
        try:
            out = _extract_text_from_txt(p)
        finally:
            p.unlink(missing_ok=True)
        self.assertIn("Xin chào", out)

    def test_extract_text_from_txt_without_chardet(self):
        import pathlib
        import tempfile as _tf

        from .readers import _extract_text_from_txt

        with _tf.NamedTemporaryFile("wb", suffix=".txt", delete=False) as f:
            f.write(b"plain ascii text")
            p = pathlib.Path(f.name)
        try:
            with patch.dict("sys.modules", {"chardet": None}):
                out = _extract_text_from_txt(p)
        finally:
            p.unlink(missing_ok=True)
        self.assertEqual(out, "plain ascii text")

    def test_extract_docx_missing_lib(self):
        from .readers import _extract_text_from_docx

        with patch.dict("sys.modules", {"docx": None}):
            with self.assertRaises(ValueError) as ctx:
                _extract_text_from_docx("x.docx")
        self.assertIn("python-docx", str(ctx.exception))

    def test_extract_pdf_missing_lib(self):
        from .readers import _extract_text_from_pdf

        with patch.dict("sys.modules", {"fitz": None}):
            with self.assertRaises(ValueError) as ctx:
                _extract_text_from_pdf("x.pdf")
        self.assertIn("PyMuPDF", str(ctx.exception))

    def test_extract_docx_success(self):
        import pathlib
        import tempfile as _tf

        from docx import Document as DocxDocument

        from .readers import _extract_text_from_docx

        with _tf.TemporaryDirectory() as d:
            p = pathlib.Path(d) / "a.docx"
            doc = DocxDocument()
            doc.add_paragraph("First paragraph of the docx file.")
            doc.add_paragraph("Second paragraph content.")
            doc.save(p)
            out = _extract_text_from_docx(p)
        self.assertIn("First paragraph", out)
        self.assertIn("Second paragraph", out)

    @patch("summaries.readers._get_http_session")
    @patch("summaries.readers._resolve_and_validate")
    def test_redirects_then_succeeds(self, _resolve, mock_get_session):
        sess = mock_get_session.return_value
        r1, r2 = MagicMock(), MagicMock()
        r1.status_code = 301
        r1.headers = {"Location": "https://final.example.com/page"}
        r1.text = ""
        r2.status_code = 200
        r2.headers = {"Content-Type": "text/html; charset=utf-8"}
        r2.text = "<html><body><p>Final page content here.</p></body></html>"
        sess.get.side_effect = [r1, r2]

        out = extract_text("https://start.example.com/old")
        self.assertIn("Final page content", out)

    @patch("summaries.readers._get_http_session")
    @patch("summaries.readers._resolve_and_validate")
    def test_http_error_raised(self, _resolve, mock_get_session):
        sess = mock_get_session.return_value
        resp = MagicMock()
        resp.status_code = 404
        resp.text = "nope"
        resp.headers = {}
        sess.get.return_value = resp

        with self.assertRaises(ValueError) as ctx:
            extract_text("https://example.com/missing")
        self.assertIn("404", str(ctx.exception))

    @patch("summaries.readers._get_http_session")
    @patch("summaries.readers._resolve_and_validate")
    def test_response_none_raises(self, _resolve, mock_get_session):
        mock_get_session.return_value.get.return_value = None

        with self.assertRaises(ValueError) as ctx:
            extract_text("https://example.com/empty")
        self.assertIn("Không thể tải URL", str(ctx.exception))

    @patch("summaries.readers._get_http_session")
    @patch("summaries.readers._resolve_and_validate")
    def test_plain_text_content_type(self, _resolve, mock_get_session):
        sess = mock_get_session.return_value
        resp = MagicMock()
        resp.status_code = 200
        resp.headers = {"Content-Type": "text/plain; charset=utf-8"}
        resp.text = "Just some plain text document."
        sess.get.return_value = resp

        out = extract_text("https://example.com/notes.txt")
        self.assertEqual(out, "Just some plain text document.")

    @patch("summaries.readers._get_http_session")
    @patch("summaries.readers._resolve_and_validate")
    def test_empty_html_raises(self, _resolve, mock_get_session):
        sess = mock_get_session.return_value
        resp = MagicMock()
        resp.status_code = 200
        resp.headers = {"Content-Type": "text/html; charset=utf-8"}
        resp.text = "<html><body></body></html>"
        sess.get.return_value = resp

        with self.assertRaises(ValueError):
            extract_text("https://example.com/empty")


class ReaderCoverageTests(TestCase):
    @patch("summaries.readers.socket.getaddrinfo")
    def test_resolve_and_validate_gaierror(self, mock_addr):
        import socket as _socket

        mock_addr.side_effect = _socket.gaierror("no such host")
        with self.assertRaises(ValueError) as ctx:
            _resolve_and_validate("definitely-not-a-host.invalid")
        self.assertIn("phân giải", str(ctx.exception))

    @patch("summaries.readers.socket.getaddrinfo")
    def test_resolve_and_validate_blocks_private_ip(self, mock_addr):
        mock_addr.return_value = [
            (2, 1, 6, "", ("10.0.0.5", 80)),
        ]
        with self.assertRaises(ValueError) as ctx:
            _resolve_and_validate("example.com")
        self.assertIn("nội bộ", str(ctx.exception))

    def test_is_private_ip_invalid_address_returns_false(self):
        self.assertFalse(_is_private_ip("not-an-ip-address"))

    def test_get_http_session_thread_safe(self):
        from .readers import _get_http_session

        session = _get_http_session()
        self.assertIs(session, _get_http_session())
        self.assertIn("User-Agent", session.headers)
        _ = getattr(_get_http_session, "__self__", None)

    def test_extract_text_from_url_requests_missing(self):
        from .readers import extract_text_from_url

        with patch.dict("sys.modules", {"requests": None}):
            with self.assertRaises(ValueError) as ctx:
                extract_text_from_url("https://example.com")
        self.assertIn("requests", str(ctx.exception))

    def test_extract_text_from_url_bs4_missing(self):
        from .readers import extract_text_from_url

        with patch.dict("sys.modules", {"bs4": None}):
            with self.assertRaises(ValueError) as ctx:
                extract_text_from_url("https://example.com")
        self.assertIn("beautifulsoup4", str(ctx.exception))

    @patch("summaries.readers._get_http_session")
    @patch("summaries.readers._resolve_and_validate")
    def test_extract_text_from_url_connection_error(self, _resolve, mock_session):
        from requests.exceptions import ConnectionError as ConnErr

        mock_session.return_value.get.side_effect = ConnErr("refused")
        with self.assertRaises(ValueError):
            extract_text("https://example.com")

    @patch("summaries.readers._get_http_session")
    @patch("summaries.readers._resolve_and_validate")
    def test_extract_text_from_url_redirect_with_empty_location(self, _resolve, mock_session):
        """Redirect response with an empty Location stops following and raises."""
        sess = mock_session.return_value
        r = MagicMock()
        r.status_code = 302
        r.headers = {"Location": ""}
        r.text = ""
        sess.get.return_value = r
        with self.assertRaises(ValueError):
            extract_text("https://example.com")

    def test_extract_text_pdf_success(self):
        import pathlib
        import tempfile as _tf

        from .readers import _extract_text_from_pdf

        with _tf.TemporaryDirectory() as d:
            p = pathlib.Path(d) / "a.pdf"
            try:
                import fitz

                doc = fitz.open()
                page = doc.new_page()
                page.insert_text((72, 72), "Hello PDF content")
                doc.save(str(p))
                doc.close()
                out = _extract_text_from_pdf(p)
                self.assertIn("Hello PDF", out)
            except ImportError:
                self.skipTest("PyMuPDF not installed")

    def test_extract_text_from_epub_bs4_missing(self):
        from .readers import _extract_text_from_epub

        with patch.dict("sys.modules", {"bs4": None}):
            with self.assertRaises(ValueError) as ctx:
                _extract_text_from_epub("x.epub")
        self.assertIn("beautifulsoup4", str(ctx.exception))

    def test_extract_text_from_epub_ebooklib_missing(self):
        from .readers import _extract_text_from_epub

        with patch.dict("sys.modules", {"ebooklib": None}):
            with self.assertRaises(ValueError) as ctx:
                _extract_text_from_epub("x.epub")
        self.assertIn("ebooklib", str(ctx.exception))

    def test_extract_text_from_epub_success(self):
        import pathlib

        from .readers import _extract_text_from_epub

        book = MagicMock()
        item = MagicMock()
        item.get_content.return_value = (
            b"<html><body><p>Epub chapter content here.</p></body></html>"
        )
        book.get_items.return_value = [item]
        with patch("ebooklib.epub.read_epub", return_value=book):
            out = _extract_text_from_epub(pathlib.Path("book.epub"))
        self.assertIn("Epub chapter content", out)

    def test_extract_text_unsupported_extension(self):
        import pathlib
        import tempfile as _tf

        with _tf.TemporaryDirectory() as d:
            p = pathlib.Path(d) / "foo.exe"
            p.write_bytes(b"data")
            with self.assertRaises(ValueError) as ctx:
                extract_text(str(p))
        self.assertIn("không được hỗ trợ", str(ctx.exception))

    def test_extract_text_missing_file(self):
        with self.assertRaises(FileNotFoundError):
            extract_text("/tmp/definitely-missing-file.txt")
