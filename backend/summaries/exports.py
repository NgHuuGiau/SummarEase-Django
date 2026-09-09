"""Export utilities for summaries."""

from __future__ import annotations

import io

from django.conf import settings
from django.http import HttpResponse
from django.template.loader import render_to_string

try:
    from weasyprint import HTML
    HAS_WEASYPRINT = True
except ImportError:
    HAS_WEASYPRINT = False

try:
    from docx import Document as DocxDocument
    HAS_DOCX = True
except ImportError:
    HAS_DOCX = False


def export_markdown(summary) -> HttpResponse:
    """Export summary as Markdown file."""
    md_content = f"""# {summary.title}

**Phương pháp:** {summary.get_method_display()}  
**Ngôn ngữ:** {summary.language}  
**Tỉ lệ:** {summary.ratio * 100:.0f}%  
**Nguồn:** {summary.document.get_source_type_display()}  
**Ngày tạo:** {summary.created_at.strftime('%d/%m/%Y %H:%M')}

---

{summary.summary_text}

---

## Từ khóa
{', '.join(f'`{tag.name}`' for tag in summary.tags.all()) if summary.tags.exists() else 'Không có'}

## Câu gốc
{chr(10).join(f'{i}. {s.sentence_text}' for i, s in enumerate(summary.sentences.all(), 1))}
"""
    response = HttpResponse(md_content, content_type="text/markdown; charset=utf-8")
    response["Content-Disposition"] = f'attachment; filename="{summary.title[:50]}.md"'
    return response


def export_docx(summary) -> HttpResponse:
    """Export summary as DOCX file."""
    if not HAS_DOCX:
        return HttpResponse("Thiếu thư viện python-docx", status=500)

    doc = DocxDocument()
    doc.add_heading(summary.title, level=1)

    # Meta table
    table = doc.add_table(rows=5, cols=2)
    table.style = "Light Grid Accent 1"
    meta = [
        ("Phương pháp", summary.get_method_display()),
        ("Ngôn ngữ", summary.language),
        ("Tỉ lệ", f"{summary.ratio * 100:.0f}%"),
        ("Nguồn", summary.document.get_source_type_display()),
        ("Ngày tạo", summary.created_at.strftime("%d/%m/%Y %H:%M")),
    ]
    for i, (label, value) in enumerate(meta):
        table.rows[i].cells[0].text = label
        table.rows[i].cells[1].text = value

    doc.add_paragraph()
    doc.add_heading("Nội dung tóm tắt", level=2)
    doc.add_paragraph(summary.summary_text)

    if summary.tags.exists():
        doc.add_paragraph()
        doc.add_heading("Từ khóa", level=2)
        keywords = ", ".join(tag.name for tag in summary.tags.all())
        doc.add_paragraph(keywords)

    if summary.sentences.exists():
        doc.add_paragraph()
        doc.add_heading("Các câu tóm tắt", level=2)
        for i, s in enumerate(summary.sentences.all(), 1):
            doc.add_paragraph(f"{i}. {s.sentence_text}", style="List Number")

    buffer = io.BytesIO()
    doc.save(buffer)
    buffer.seek(0)

    response = HttpResponse(
        buffer.getvalue(),
        content_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    )
    response["Content-Disposition"] = f'attachment; filename="{summary.title[:50]}.docx"'
    return response


def export_pdf(summary) -> HttpResponse:
    """Export summary as PDF file using WeasyPrint."""
    if not HAS_WEASYPRINT:
        return HttpResponse("Thiếu thư viện weasyprint", status=500)

    html_content = render_to_string("summaries/export_pdf.html", {
        "summary": summary,
        "tags": summary.tags.all(),
        "sentences": summary.sentences.all(),
    })

    pdf_file = io.BytesIO()
    HTML(string=html_content, base_url=settings.STATIC_ROOT).write_pdf(pdf_file)
    pdf_file.seek(0)

    response = HttpResponse(pdf_file.getvalue(), content_type="application/pdf")
    response["Content-Disposition"] = f'attachment; filename="{summary.title[:50]}.pdf"'
    return response


def export_summary(request, pk: int, format: str):
    """Main export view - routes to appropriate format."""
    from .models import Summary
    from django.shortcuts import get_object_or_404

    summary = get_object_or_404(
        Summary.objects.select_related("document").prefetch_related("tags", "sentences"),
        pk=pk,
        user=request.user,
    )

    if format == "md":
        return export_markdown(summary)
    elif format == "docx":
        return export_docx(summary)
    elif format == "pdf":
        return export_pdf(summary)
    else:
        return HttpResponse("Định dạng không hỗ trợ", status=400)